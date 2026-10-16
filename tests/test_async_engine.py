import asyncio
import base64
import json
import threading
from io import BytesIO
from unittest.mock import MagicMock
from zipfile import ZipFile

import httpx
import pytest
from docx import Document

from app import engine as engine_module
from app.blueprint import BLUEPRINT_PATH, build_solution_blueprint
from app.engine import BaseAgent, CodingAgent, DocReviewerAgent, OrchestratorEngine
from app.main import create_app
from app.schemas import ApprovalRequest, CreateRunRequest, FileDiff, RunStatus, TaskType


async def test_blueprint_document_is_proposed_and_written_only_on_approval(tmp_path, monkeypatch):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    engine = OrchestratorEngine(workspace_root=tmp_path, write_on_approval=True)
    run = engine.create_run(CreateRunRequest(prompt="Build a small API."))
    await engine.run_pipeline(run.run_id)
    assert run.status is RunStatus.AWAITING_APPROVAL
    document_path = "docs/Solution_Blueprint.docx"
    assert document_path in {diff.path for diff in run.diffs}
    assert not (tmp_path / document_path).exists()
    json.dumps(run.model_dump(mode="json"))
    engine.submit_approval(
        run.run_id, run.pending_checkpoint.checkpoint_id,
        ApprovalRequest(decision="approved"),
    )
    with ZipFile(tmp_path / document_path) as archive:
        document = archive.read("word/document.xml").decode()
    for heading in (
        "Solution Blueprint", "System Architecture Overview", "Component &amp; API Specs",
        "Data Models", "File Structure",
    ):
        assert heading in document
    payload = run.context_diffs["blueprint_documents"][document_path]
    assert payload["encoding"] == "base64"
    assert base64.b64decode(payload["content"], validate=True) == (tmp_path / document_path).read_bytes()
    assert Document(tmp_path / document_path).core_properties.title == "Solution Blueprint"


@pytest.mark.parametrize("source_format", ["mapping", "json", "markdown"])
def test_blueprint_document_formats_sections_tables_and_code(source_format):
    details = {
        "architecture": "## Service Boundary\n\n- API service\n- Database\n",
        "component_api_specs": "**Users API** accepts `GET /users`.\n\n```http\nGET /users\n```",
        "data_models": "```python\nclass User:\n    id: int\n```",
        "files": [{"path": "app/users.py", "purpose": "User endpoint"}],
        "dependencies": ["FastAPI"],
    }
    source = details
    if source_format == "json":
        source = json.dumps(details)
    elif source_format == "markdown":
        source = (
            f"# Solution Blueprint\n\n## System Architecture Overview\n{details['architecture']}"
            f"\n## Component & API Specs\n{details['component_api_specs']}"
            f"\n## Data Models\n{details['data_models']}"
            "\n## File Structure\n\n| File Path | Purpose |\n| --- | --- |\n"
            "| app/users.py | User endpoint |\n"
        )
    document = Document(BytesIO(build_solution_blueprint(source)))
    paragraphs = document.paragraphs
    assert paragraphs[0].text == "Solution Blueprint"
    assert paragraphs[0].style.name == "Title"
    assert any(item.text == "API service" and item.style.name == "List Bullet" for item in paragraphs)
    assert any(item.text == "GET /users" and item.style.name == "Blueprint Code" for item in paragraphs)
    assert any(item.text.startswith("class User:") and item.style.name == "Blueprint Code" for item in paragraphs)
    assert any(run.bold and run.text == "Users API" for item in paragraphs for run in item.runs)
    assert any(run.text == "GET /users" and run.font.name == "Consolas" for item in paragraphs for run in item.runs)
    assert len(document.tables) == 1
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == [
        ["File Path", "Purpose"], ["app/users.py", "User endpoint"],
    ]


@pytest.mark.parametrize("write_enabled,decision", [(True, "rejected"), (False, "approved")])
async def test_blueprint_proposal_without_approved_writes(tmp_path, monkeypatch, write_enabled, decision):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    engine = OrchestratorEngine(workspace_root=tmp_path, write_on_approval=write_enabled)
    run = engine.create_run(CreateRunRequest(prompt="Plan an API."))
    await engine.run_pipeline(run.run_id)
    assert BLUEPRINT_PATH in {diff.path for diff in run.diffs}
    engine.submit_approval(
        run.run_id, run.pending_checkpoint.checkpoint_id, ApprovalRequest(decision=decision),
    )
    assert list(tmp_path.iterdir()) == []


async def test_existing_blueprint_detects_conflicts_and_can_be_replaced(tmp_path, monkeypatch):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    destination = tmp_path / BLUEPRINT_PATH
    destination.parent.mkdir()
    original = build_solution_blueprint({"architecture": "Previous architecture"})
    destination.write_bytes(original)
    engine = OrchestratorEngine(workspace_root=tmp_path, write_on_approval=True)
    run = engine.create_run(CreateRunRequest(prompt="Revise the API."))
    await engine.run_pipeline(run.run_id)
    assert next(diff for diff in run.diffs if diff.path == BLUEPRINT_PATH).status == "modified"
    assert destination.read_bytes() == original
    checkpoint = run.pending_checkpoint.checkpoint_id
    destination.write_bytes(b"An external edit")
    with pytest.raises(ValueError, match="Workspace changed since review"):
        engine.submit_approval(run.run_id, checkpoint, ApprovalRequest(decision="approved"))
    assert destination.read_bytes() == b"An external edit"
    assert not (tmp_path / "app").exists()
    assert run.status is RunStatus.AWAITING_APPROVAL
    destination.write_bytes(original)
    engine.submit_approval(run.run_id, checkpoint, ApprovalRequest(decision="approved"))
    assert Document(destination).paragraphs[0].text == "Solution Blueprint"


@pytest.mark.parametrize("before_planning", [True, False])
async def test_blueprint_rejects_symlinked_directory(tmp_path, monkeypatch, before_planning):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    outside = tmp_path / "outside"
    outside.mkdir()
    if before_planning:
        (workspace / "docs").symlink_to(outside, target_is_directory=True)
    engine = OrchestratorEngine(workspace_root=workspace, write_on_approval=True)
    run = engine.create_run(CreateRunRequest(prompt="Plan an API."))
    await engine.run_pipeline(run.run_id)
    if before_planning:
        assert run.status is RunStatus.FAILED
        assert "Symlinks are not allowed" in run.state_history[-1].output_summary
    else:
        (workspace / "docs").symlink_to(outside, target_is_directory=True)
        with pytest.raises(ValueError, match="Symlinks are not allowed"):
            engine.submit_approval(
                run.run_id, run.pending_checkpoint.checkpoint_id,
                ApprovalRequest(decision="approved"),
            )
    assert list(outside.iterdir()) == []


def test_binary_workspace_writes_roll_back_on_failure(tmp_path, monkeypatch):
    destination = tmp_path / "existing.docx"
    original = build_solution_blueprint({"architecture": "Original"})
    destination.write_bytes(original)
    replace = engine_module.os.replace

    def fail_source_replace(source, target):
        if target.name == "new.py":
            raise OSError("Simulated write failure")
        return replace(source, target)

    monkeypatch.setattr(engine_module.os, "replace", fail_source_replace)
    with pytest.raises(OSError, match="Simulated write failure"):
        engine_module._write_workspace(
            tmp_path, {"existing.docx": b"replacement", "new.py": "VALUE = 1\n"},
            {"existing.docx": original},
        )
    assert destination.read_bytes() == original
    assert list(tmp_path.iterdir()) == [destination]


@pytest.mark.parametrize(
    "mode,names",
    [
        (TaskType.CODE, ["Coder"]),
        (TaskType.REVIEW_DOCS, ["DocReviewer"]),
        (TaskType.PIPELINE, ["Coder", "DocReviewer"]),
    ],
)
async def test_mode_specific_pipelines(mode, names):
    engine = OrchestratorEngine()
    run = engine.create_run(CreateRunRequest(prompt="Add pagination.", pipeline_mode=mode))
    assert run.status is RunStatus.QUEUED
    assert await engine.run_pipeline(run.run_id) is run
    assert run.status is RunStatus.AWAITING_APPROVAL
    assert [step.agent_name for step in run.state_history] == names
    assert len(run.diffs) == len(names)
    assert all(isinstance(diff, FileDiff) for diff in run.diffs)
    assert run.pending_checkpoint is not None


async def test_agents_return_models_and_structured_findings():
    generated = await CodingAgent().execute("Add pagination.", target_files=["app/users.py"])
    assert isinstance(generated["diffs"][0], FileDiff)
    reviewed = await DocReviewerAgent().execute(
        "Document pagination.",
        json.dumps(
            {
                "diffs": [diff.model_dump(mode="json") for diff in generated["diffs"]],
            }
        ),
    )
    assert reviewed["findings"][0]["path"] == "app/users.py"
    assert isinstance(reviewed["diffs"][0], FileDiff)


async def test_context_forwarding_and_multiple_revision_rounds():
    class RecordingReviewer(DocReviewerAgent):
        async def execute(self, prompt, context="", target_files=None):
            self.context = json.loads(context)
            return await super().execute(prompt, context, target_files)

    reviewer = RecordingReviewer()
    engine = OrchestratorEngine(doc_reviewer_agent=reviewer)
    run = engine.create_run(CreateRunRequest(prompt="Add pagination."))
    await engine.run_pipeline(run.run_id)
    assert reviewer.context["code_outputs"]
    assert reviewer.context["diffs"]
    assert run.context_diffs["DocReviewer"]["findings"]
    for feedback in ["Limit to 100.", "Document errors."]:
        old_checkpoint = run.pending_checkpoint.checkpoint_id
        engine.submit_approval(
            run.run_id,
            old_checkpoint,
            ApprovalRequest(
                decision="request_changes",
                feedback=feedback,
            ),
        )
        assert run.status is RunStatus.RUNNING
        assert feedback in run.context_diffs["feedback"]
        await engine.run_pipeline(run.run_id)
        assert run.pending_checkpoint.checkpoint_id != old_checkpoint
        assert len(run.diffs) == 2
        assert feedback in run.diffs[0].patch
    assert reviewer.context["feedback"] == ["Limit to 100.", "Document errors."]
    assert len(run.state_history) == 6


async def test_duplicate_execution_is_serialized():
    class YieldingCoder(CodingAgent):
        async def execute(self, prompt, context="", target_files=None):
            await asyncio.sleep(0)
            return await super().execute(prompt, context, target_files)

    engine = OrchestratorEngine(coding_agent=YieldingCoder())
    run = engine.create_run(CreateRunRequest(prompt="Test"))
    await asyncio.gather(engine.run_pipeline(run.run_id), engine.run_pipeline(run.run_id))
    assert len(run.state_history) == 2


async def test_cancelled_agent_marks_run_failed():
    class CancelledAgent(BaseAgent):
        agent_name = "Cancelled"

        async def execute(self, prompt, context="", target_files=None):
            raise asyncio.CancelledError

    engine = OrchestratorEngine(coding_agent=CancelledAgent())
    run = engine.create_run(CreateRunRequest(prompt="Test"))
    with pytest.raises(asyncio.CancelledError):
        await engine.run_pipeline(run.run_id)
    assert run.status is RunStatus.FAILED
    assert run.current_step is None
    assert run.state_history[-1].status == "failed"


async def test_api_reads_running_state_while_agent_waits():
    started = asyncio.Event()
    release = asyncio.Event()

    class WaitingCoder(CodingAgent):
        async def execute(self, prompt, context="", target_files=None):
            started.set()
            await release.wait()
            return await super().execute(prompt, context, target_files)

    engine = OrchestratorEngine(coding_agent=WaitingCoder())
    run = engine.create_run(CreateRunRequest(prompt="Test"))
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=create_app(engine)),
        base_url="http://test",
    ) as client:
        task = asyncio.create_task(engine.run_pipeline(run.run_id))
        try:
            await asyncio.wait_for(started.wait(), timeout=1)
            response = await asyncio.wait_for(client.get(f"/v1/runs/{run.run_id}"), timeout=1)
            assert response.json()["status"] == "running"
            assert response.json()["current_step"] == "Coder"
        finally:
            release.set()
            await asyncio.wait_for(task, timeout=1)


@pytest.mark.parametrize("failure", ["timeout", "unavailable", "empty_response"])
async def test_gemini_waits_do_not_block_api_and_fallback_leaves_running(monkeypatch, failure):
    monkeypatch.setenv("GEMINI_API_KEY", "test-key")
    monkeypatch.setattr(engine_module, "_GEMINI_ATTEMPT_TIMEOUT_SECONDS", 0.05)
    started = threading.Event()
    finished = threading.Event()
    release = threading.Event()
    event_loop_thread = threading.get_ident()
    original_sleep = asyncio.sleep

    def generate(**kwargs):
        assert threading.get_ident() != event_loop_thread
        started.set()
        if failure == "timeout":
            try:
                if not release.wait(timeout=2):
                    raise RuntimeError("Test worker was not released")
            finally:
                finished.set()
        elif failure == "unavailable":
            raise RuntimeError("503 UNAVAILABLE")
        return None

    async def backoff(delay):
        await original_sleep(0)

    client_mock = MagicMock()
    client_mock.models.generate_content = MagicMock(side_effect=generate)
    monkeypatch.setattr(engine_module.genai, "Client", MagicMock(return_value=client_mock))
    monkeypatch.setattr(engine_module.asyncio, "sleep", backoff)
    engine = OrchestratorEngine()
    run = engine.create_run(CreateRunRequest(prompt="Test", pipeline_mode=TaskType.CODE))
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=create_app(engine)), base_url="http://test"
    ) as client:
        task = asyncio.create_task(engine.run_pipeline(run.run_id))
        try:
            assert await asyncio.to_thread(started.wait, 1)
            response = await asyncio.wait_for(client.get(f"/v1/runs/{run.run_id}"), timeout=1)
            assert response.status_code == 200
            await asyncio.wait_for(task, timeout=1)
        finally:
            release.set()
            if failure == "timeout":
                assert await asyncio.to_thread(finished.wait, 1)
            if not task.done():
                task.cancel()
                await asyncio.gather(task, return_exceptions=True)
    assert run.status is RunStatus.AWAITING_APPROVAL
    assert run.current_step is None
    assert "mock" in run.state_history[-1].output_summary
    assert client_mock.models.generate_content.call_count == (4 if failure == "unavailable" else 1)
