"""Render structured or Markdown solution blueprints as Word documents."""

import json
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from markdown_it import MarkdownIt

BLUEPRINT_PATH = "docs/Solution_Blueprint.docx"
_MARKDOWN = MarkdownIt().enable("table")
_SECTIONS = (
    ("System Architecture Overview", "architecture"),
    ("Component & API Specs", "component_api_specs"),
    ("Data Models", "data_models"),
    ("File Structure", "files"),
)


def _markdown_value(value: Any) -> str:
    if isinstance(value, str):
        return value
    if isinstance(value, dict):
        return "\n\n".join(
            f"**{key}**\n\n{_markdown_value(item)}" for key, item in value.items()
        )
    if isinstance(value, list):
        return "\n\n".join(_markdown_value(item) for item in value)
    return json.dumps(value, ensure_ascii=True)


def _parse_blueprint(blueprint: dict[str, Any] | str) -> dict[str, Any]:
    if isinstance(blueprint, dict):
        return blueprint
    try:
        parsed = json.loads(blueprint)
    except json.JSONDecodeError:
        parsed = None
    if isinstance(parsed, dict):
        return parsed
    lines = blueprint.splitlines()
    headings = []
    tokens = _MARKDOWN.parse(blueprint)
    for index, token in enumerate(tokens):
        if token.type == "heading_open" and token.map:
            title = tokens[index + 1].content.casefold()
            for label, key in _SECTIONS:
                if title in {label.casefold(), key.replace("_", " ")}:
                    headings.append((key, token.map[0], token.map[1]))
    if not headings:
        return {"architecture": blueprint}
    sections = {}
    for index, (key, start, end) in enumerate(headings):
        stop = headings[index + 1][1] if index + 1 < len(headings) else len(lines)
        sections[key] = "\n".join(lines[end:stop])
    return sections


def _inline(paragraph: Any, tokens: list[Any]) -> None:
    bold = False
    italic = False
    for token in tokens:
        if token.type == "strong_open":
            bold = True
        elif token.type == "strong_close":
            bold = False
        elif token.type == "em_open":
            italic = True
        elif token.type == "em_close":
            italic = False
        elif token.type in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()
        elif token.type in {"text", "code_inline"}:
            run = paragraph.add_run(token.content)
            run.bold = bold
            run.italic = italic
            if token.type == "code_inline":
                run.font.name = "Consolas"
                run.font.size = Pt(10)


def _render_markdown(document: Any, text: str) -> None:
    paragraph = None
    lists = []
    table = None
    row = None
    cell_index = 0
    for token in _MARKDOWN.parse(text):
        if token.type in {"bullet_list_open", "ordered_list_open"}:
            lists.append(token.type)
        elif token.type in {"bullet_list_close", "ordered_list_close"}:
            lists.pop()
        elif token.type == "heading_open":
            paragraph = document.add_heading(level=min(int(token.tag[1:]) + 1, 4))
        elif token.type == "paragraph_open":
            style = None
            if lists:
                style = "List Number" if lists[-1] == "ordered_list_open" else "List Bullet"
            paragraph = document.add_paragraph(style=style)
        elif token.type in {"fence", "code_block"}:
            document.add_paragraph(token.content.rstrip("\n"), style="Blueprint Code")
        elif token.type == "table_open":
            table = document.add_table(rows=0, cols=0)
            table.style = "Light Shading Accent 1"
        elif token.type == "tr_open":
            row = None
            cell_index = 0
        elif token.type in {"th_open", "td_open"} and table is not None:
            if row is None:
                row = table.add_row()
            if cell_index >= len(table.columns):
                table.add_column(Inches(2))
            paragraph = row.cells[cell_index].paragraphs[0]
            cell_index += 1
        elif token.type == "inline" and paragraph is not None:
            _inline(paragraph, token.children or [])


def build_solution_blueprint(blueprint: dict[str, Any] | str) -> bytes:
    """Build a styled DOCX in memory, without writing workspace files."""
    details = _parse_blueprint(blueprint)
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        document.styles[name].font.color.rgb = RGBColor.from_string("174A5B")
    code = document.styles.add_style("Blueprint Code", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code.font.size = Pt(9)
    code.paragraph_format.left_indent = Inches(0.15)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F0F3F5")
    code.element.get_or_add_pPr().append(shading)
    document.core_properties.title = "Solution Blueprint"
    document.add_heading("Solution Blueprint", 0)
    document.add_paragraph("Architecture and Implementation Specification", style="Subtitle")
    for label, key in _SECTIONS:
        document.add_heading(label, 1)
        value = details.get(key)
        if key == "files" and isinstance(value, list):
            table = document.add_table(rows=1, cols=2)
            table.style = "Light Shading Accent 1"
            table.rows[0].cells[0].text = "File Path"
            table.rows[0].cells[1].text = "Purpose"
            for file in value:
                cells = table.add_row().cells
                cells[0].text = str(file.get("path", ""))
                cells[1].text = str(file.get("purpose", ""))
        else:
            _render_markdown(document, _markdown_value(value) if value else "Not specified.")
    if details.get("dependencies"):
        document.add_heading("Dependencies", 1)
        dependencies = details["dependencies"]
        if isinstance(dependencies, list):
            for dependency in dependencies:
                document.add_paragraph(str(dependency), style="List Bullet")
        else:
            _render_markdown(document, str(dependencies))
    section.footer.paragraphs[0].text = "Solution Blueprint | Architecture Specification"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()